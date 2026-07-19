from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any

SCHEME_CODE_RE = re.compile(r"^[0-9]{5,6}$")
OLE2_MAGIC = b"\xd0\xcf\x11\xe0"
ZIP_MAGIC = b"PK"


def normalize_scheme_name(name: str) -> str:
    cleaned = name.lower()
    cleaned = re.sub(r"\s+", " ", cleaned)
    for token in (" - regular plan", " - direct plan", " regular plan", " direct plan"):
        cleaned = cleaned.replace(token, "")
    cleaned = re.sub(r"[^a-z0-9 ]+", "", cleaned)
    return cleaned.strip()


def parse_ter_rows(rows: list[dict]) -> list[dict[str, Any]]:
    parsed: list[dict[str, Any]] = []
    for row in rows:
        scheme_name = str(row.get("Scheme_Name") or row.get("scheme_name") or "").strip()
        if not scheme_name:
            continue
        as_of = _parse_ter_date(row.get("TER_Date"))
        regular_ter = _parse_decimal(row.get("R_TER"))
        if regular_ter is not None:
            parsed.append(
                {
                    "scheme_name": scheme_name,
                    "plan_type": "REGULAR",
                    "ter_percent": regular_ter,
                    "as_of_date": as_of,
                }
            )
    return parsed


def parse_ter_tracker_csv(body: bytes) -> list[dict[str, Any]]:
    text = body.decode("utf-8", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        return []
    fields = {name.strip().lower(): name for name in reader.fieldnames}
    name_col = next((fields[k] for k in fields if "scheme name" in k), None)
    regular_col = next(
        (fields[k] for k in fields if "regular plan" in k and "total ter" in k),
        None,
    )
    if not name_col or not regular_col:
        return []

    as_of = date.today().replace(day=1)
    parsed: list[dict[str, Any]] = []
    for row in reader:
        scheme_name = (row.get(name_col) or "").strip()
        ter = _parse_decimal(row.get(regular_col))
        if scheme_name and ter is not None:
            parsed.append(
                {
                    "scheme_name": scheme_name,
                    "plan_type": "REGULAR",
                    "ter_percent": ter,
                    "as_of_date": as_of,
                }
            )
    return parsed


def parse_aaum_json_blocks(blocks: list[dict]) -> list[dict[str, Any]]:
    parsed: list[dict[str, Any]] = []
    for block in blocks:
        for scheme in block.get("schemes") or []:
            if not isinstance(scheme, dict):
                continue
            scheme_code = str(scheme.get("AMFI_Code") or scheme.get("amfi_code") or "").strip()
            if not SCHEME_CODE_RE.match(scheme_code):
                continue
            aum_block = scheme.get("AverageAumForTheMonth") or {}
            if not isinstance(aum_block, dict):
                continue
            crores = _parse_decimal(
                aum_block.get("ExcludingFundOfFundsDomesticButIncludingFundOfFundsOverseas")
            )
            if crores is None:
                continue
            parsed.append(
                {
                    "scheme_code": scheme_code,
                    "scheme_name": str(scheme.get("SchemeNAVName") or "").strip(),
                    "aum_crores": crores,
                }
            )
    return parsed


def parse_aum_workbook(body: bytes, *, as_of_date: date) -> list[dict[str, Any]]:
    if body.startswith(OLE2_MAGIC):
        return _parse_xls(body, as_of_date=as_of_date)
    if body.startswith(ZIP_MAGIC):
        return _parse_xlsx(body, as_of_date=as_of_date)
    text = body.decode("utf-8", errors="replace")
    if "<table" in text.lower():
        return _parse_html_table(text, as_of_date=as_of_date)
    return _parse_delimited(text, as_of_date=as_of_date)


def _parse_xls(body: bytes, *, as_of_date: date) -> list[dict[str, Any]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=body)
    records: list[dict[str, Any]] = []
    for sheet in book.sheets():
        records.extend(_parse_tabular_rows(_sheet_to_rows(sheet), as_of_date=as_of_date))
    return records


def _parse_xlsx(body: bytes, *, as_of_date: date) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        rows = [[cell.value for cell in row] for row in sheet.iter_rows(values_only=False)]
        flat_rows = []
        for row in rows:
            flat_rows.append(["" if value is None else str(value).strip() for value in row])
        records.extend(_parse_tabular_rows(flat_rows, as_of_date=as_of_date))
    return records


def _sheet_to_rows(sheet) -> list[list[str]]:
    rows: list[list[str]] = []
    for r in range(sheet.nrows):
        row: list[str] = []
        for c in range(sheet.ncols):
            value = sheet.cell_value(r, c)
            if isinstance(value, float) and value.is_integer():
                row.append(str(int(value)))
            else:
                row.append(str(value).strip())
        rows.append(row)
    return rows


def _parse_tabular_rows(rows: list[list[str]], *, as_of_date: date) -> list[dict[str, Any]]:
    if len(rows) < 2:
        return []
    header_idx = _find_header_row(rows)
    header = [cell.lower() for cell in rows[header_idx]]
    scheme_idx = _find_col(header, "scheme code", "code")
    aum_idx = _find_col(header, "aum", "net assets", "average aum")
    if scheme_idx < 0:
        scheme_idx = 0
    if aum_idx < 0:
        aum_idx = min(scheme_idx + 1, len(header) - 1)

    parsed: list[dict[str, Any]] = []
    for row in rows[header_idx + 1 :]:
        if scheme_idx >= len(row):
            continue
        scheme_code = row[scheme_idx].strip()
        if not SCHEME_CODE_RE.match(scheme_code):
            continue
        aum = _parse_decimal(row[aum_idx] if aum_idx < len(row) else None)
        if aum is None:
            continue
        parsed.append({"scheme_code": scheme_code, "aum_crores": aum, "as_of_date": as_of_date})
    return parsed


def _parse_delimited(text: str, *, as_of_date: date) -> list[dict[str, Any]]:
    delimiter = "\t" if "\t" in text else ","
    rows = [line.split(delimiter) for line in text.splitlines() if line.strip()]
    return _parse_tabular_rows(rows, as_of_date=as_of_date)


def _parse_html_table(html: str, *, as_of_date: date) -> list[dict[str, Any]]:
    row_pattern = re.compile(r"<tr[^>]*>(.*?)</tr>", re.IGNORECASE | re.DOTALL)
    cell_pattern = re.compile(r"<t[dh][^>]*>(.*?)</t[dh]>", re.IGNORECASE | re.DOTALL)
    tag_pattern = re.compile(r"<[^>]+>")
    rows: list[list[str]] = []
    for row_match in row_pattern.finditer(html):
        cells = [
            tag_pattern.sub("", cell).strip()
            for cell in cell_pattern.findall(row_match.group(1))
        ]
        if cells:
            rows.append(cells)
    return _parse_tabular_rows(rows, as_of_date=as_of_date)


def _find_header_row(rows: list[list[str]]) -> int:
    for idx, row in enumerate(rows[:40]):
        joined = " ".join(cell.lower() for cell in row)
        if "scheme code" in joined or "scheme_code" in joined:
            return idx
    return 0


def _find_col(header: list[str], *needles: str) -> int:
    for idx, cell in enumerate(header):
        for needle in needles:
            if needle in cell:
                return idx
    return -1


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None:
        return None
    text = str(value).strip().replace(",", "").replace("%", "")
    if not text or text.lower() in {"-", "na", "n.a.", "n/a"}:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _parse_ter_date(value: Any) -> date:
    if not value:
        return date.today().replace(day=1)
    text = str(value)
    for fmt in ("%Y-%m-%dT%H:%M:%S.%fZ", "%Y-%m-%d", "%d-%b-%Y"):
        try:
            if "T" in fmt:
                return datetime.strptime(text[:23], fmt).date()
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return date.today().replace(day=1)


def crores_to_inr(crores: Decimal) -> Decimal:
    return (crores * Decimal("10000000")).quantize(Decimal("0.01"))


def quarter_end_from_period_label(period_label: str) -> date:
    match = re.search(
        r"(?i)(January|February|March|April|May|June|July|August|September|October|November|December)"
        r"\s*-\s*"
        r"(January|February|March|April|May|June|July|August|September|October|November|December)"
        r"\s*(\d{4})",
        period_label.strip(),
    )
    if not match:
        raise ValueError(f"Cannot parse AMFI AAUM period label: {period_label}")
    end_month_name = match.group(2).title()
    year = int(match.group(3))
    month_number = datetime.strptime(end_month_name, "%B").month
    if month_number == 12:
        next_month = date(year + 1, 1, 1)
    else:
        next_month = date(year, month_number + 1, 1)
    return next_month - timedelta(days=1)
